#!/usr/bin/env python3
"""Export a validated Equinet A1 review package to JSON, Markdown, CSV and Excel."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SKILL_DIR = Path(__file__).resolve().parents[1]
PROFILE_DIR = Path(__file__).resolve().parents[3]
REVIEW_SKILL = PROFILE_DIR / "skills" / "ranked-prospect-review-package"
REVIEW_SCHEMA = REVIEW_SKILL / "references" / "review-package.schema.json"
MANIFEST_SCHEMA = SKILL_DIR / "references" / "export-manifest.schema.json"
REVIEW_SCRIPT_DIR = REVIEW_SKILL / "scripts"
if str(REVIEW_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(REVIEW_SCRIPT_DIR))
from build_review_package import validate_candidate  # noqa: E402

REVIEW_COLUMNS = [
    "candidate_id", "overall_rank", "segment_rank", "recommended_review_status", "reviewer_decision",
    "run_id", "discovered_at", "segment", "prospect_type", "display_name", "person_name", "role_title",
    "organisation_name", "website", "domain", "country_code", "state_region", "city", "postal_code", "full_address", "geography_status",
    "public_emails", "public_phones", "public_profiles", "public_contact_details", "icp_score", "score_band", "score_model_version",
    "score_rationale", "scoring_outcome", "applied_band_rules", "confirmed_criteria", "uncertain_criteria",
    "confidence_level", "confidence_score", "confidence_limitations", "minimum_data_status",
    "missing_minimum_fields", "exclusion_status", "batch_duplicate_status", "twenty_duplicate_status",
    "hubspot_duplicate_status", "source_urls", "evidence_summary", "next_action", "priority_rank",
    "suggested_territory", "workflow_stage", "canonical_review_decision", "integration_limitations",
    "a2_handoff_status", "recommendation_rationale", "schema_version",
]
QUEUE_COLUMNS = [
    "overall_rank", "recommended_review_status", "candidate_id", "display_name", "person_name", "role_title", "segment", "prospect_type",
    "city", "state_region", "postal_code", "full_address", "icp_score", "score_band", "confidence_score", "confidence_level",
    "scoring_outcome", "confirmed_criteria", "uncertain_criteria", "batch_duplicate_status",
    "hubspot_duplicate_status", "twenty_duplicate_status", "next_action", "public_contact_details", "source_urls", "recommendation_rationale",
]


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise ValueError(f"{path} must contain a JSON object.")
    return value


def validate_schema(path: Path, value: Any, label: str) -> None:
    schema = load_json(path)
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    errors = sorted(validator.iter_errors(value), key=lambda error: list(error.absolute_path))
    if errors:
        detail = " | ".join(f"/{'/'.join(str(part) for part in error.absolute_path)}: {error.message}" for error in errors)
        raise ValueError(f"Invalid {label}: {detail}")


def as_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def values_for_type(contacts: list[dict[str, Any]], contact_type: str) -> list[str]:
    return [entry["value"] for entry in contacts if entry["contact_type"] == contact_type]


def flatten(item: dict[str, Any]) -> dict[str, Any]:
    candidate = item["prospect_candidate"]
    identity = candidate["identity"]
    person = identity.get("person") or {}
    organisation = identity.get("organisation") or {}
    location = candidate["location"]
    contacts = candidate["public_contacts"]
    scoring = candidate["scoring"]
    confidence = candidate["confidence"]
    qualification = candidate["qualification"]
    duplicate = candidate["duplicate_check"]
    recommendation = candidate["recommendation"]
    workflow = candidate["workflow"]
    evidence_summary = [
        {
            "evidence_id": entry["evidence_id"],
            "source_name": entry["source_name"],
            "excerpt": entry["evidence_excerpt"],
            "supports_claims": entry["supports_claims"],
        }
        for entry in candidate["source_evidence"]
    ]
    other_profiles = [entry["value"] for entry in contacts if entry["contact_type"] not in {"email", "phone"}]
    return {
        "candidate_id": candidate["candidate_id"],
        "overall_rank": item["overall_rank"],
        "segment_rank": item["segment_rank"],
        "recommended_review_status": item["recommended_review_status"],
        "reviewer_decision": "",
        "run_id": candidate["run_id"],
        "discovered_at": candidate["discovered_at"],
        "segment": candidate["segment"],
        "prospect_type": candidate["prospect_type"],
        "display_name": identity["display_name"],
        "person_name": person.get("full_name"),
        "role_title": person.get("role_title"),
        "organisation_name": organisation.get("name"),
        "website": organisation.get("website"),
        "domain": organisation.get("domain"),
        "country_code": location["country_code"],
        "state_region": location.get("state_region"),
        "city": location.get("city"),
        "postal_code": location.get("postal_code"),
        "full_address": location.get("public_address"),
        "geography_status": location["geography_status"],
        "public_emails": as_json(values_for_type(contacts, "email")),
        "public_phones": as_json(values_for_type(contacts, "phone")),
        "public_profiles": as_json(other_profiles),
        "public_contact_details": as_json([
            {
                "contact_type": entry["contact_type"],
                "value": entry["value"],
                "label": entry.get("label"),
                "evidence_ids": entry["evidence_ids"],
            }
            for entry in contacts
        ]),
        "icp_score": scoring.get("score"),
        "score_band": scoring.get("band"),
        "score_model_version": scoring.get("model_version"),
        "score_rationale": scoring.get("rationale") or scoring.get("block_reason"),
        "scoring_outcome": item["scoring_outcome"],
        "applied_band_rules": as_json(item["applied_band_rules"]),
        "confirmed_criteria": as_json(item["confirmed_criteria"]),
        "uncertain_criteria": as_json(item["uncertain_criteria"]),
        "confidence_level": confidence["level"],
        "confidence_score": confidence["score"],
        "confidence_limitations": as_json(confidence["limitations"]),
        "minimum_data_status": qualification["minimum_data_status"],
        "missing_minimum_fields": as_json(qualification["missing_minimum_fields"]),
        "exclusion_status": qualification["exclusion_status"],
        "batch_duplicate_status": duplicate["batch"]["status"],
        "twenty_duplicate_status": duplicate["twenty"]["status"],
        "hubspot_duplicate_status": duplicate["hubspot"]["status"],
        "source_urls": as_json(item["source_urls"]),
        "evidence_summary": as_json(evidence_summary),
        "next_action": recommendation["next_action"],
        "priority_rank": recommendation["priority_rank"],
        "suggested_territory": recommendation.get("suggested_territory"),
        "workflow_stage": workflow["stage"],
        "canonical_review_decision": workflow["review_decision"],
        "integration_limitations": as_json(item["integration_limitations"]),
        "a2_handoff_status": item["a2_handoff"]["status"],
        "recommendation_rationale": item["recommendation_rationale"],
        "schema_version": candidate["schema_version"],
    }


def write_json(package: dict[str, Any], path: Path) -> None:
    path.write_text(json.dumps(package, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def md_escape(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")


def write_markdown(package: dict[str, Any], path: Path) -> None:
    summary = package["summary"]
    lines = [
        f"# Equinet A1 Prospect Review — {package['package_id']}",
        "",
        f"**Generated:** `{package['generated_at']}`  ",
        f"**Candidates:** `{summary['total']}` — Accept `{summary['accept']}`, Needs Research `{summary['needs_research']}`, Reject `{summary['reject']}`  ",
        "**Integrations:** HubSpot unavailable; Twenty unavailable; A2 not triggered.  ",
        "**Human decision:** Pending for every candidate.",
        "",
        "## Review queue",
        "",
        "| Rank | Candidate | Segment | ICP | Confidence | Recommendation | Human decision |",
        "|---:|---|---|---|---|---|---|",
    ]
    for item in package["candidates"]:
        candidate = item["prospect_candidate"]
        scoring = candidate["scoring"]
        confidence = candidate["confidence"]
        score_text = "Blocked" if scoring["status"] == "blocked" else f"{scoring['score']} / {scoring['band']}"
        lines.append(
            f"| {item['overall_rank']} | {md_escape(candidate['identity']['display_name'])} | {candidate['segment']} | "
            f"{score_text} | {confidence['score']} / {confidence['level']} | {item['recommended_review_status'].replace('_', ' ').title()} | Pending |"
        )
    for item in package["candidates"]:
        candidate = item["prospect_candidate"]
        lines.extend([
            "",
            f"## {item['overall_rank']}. {candidate['identity']['display_name']}",
            "",
            f"- **Candidate ID:** `{candidate['candidate_id']}`",
            f"- **Type:** `{candidate['segment']}` / `{candidate['prospect_type']}`",
            f"- **Location:** {candidate['location'].get('city') or 'Unknown'}, {candidate['location'].get('state_region') or 'Unknown'}",
            f"- **Full address:** {candidate['location'].get('public_address') or 'Not published'}",
            f"- **Primary named contact:** {(candidate['identity'].get('person') or {}).get('full_name') or 'Not identified'}",
            f"- **Primary contact role:** {(candidate['identity'].get('person') or {}).get('role_title') or 'Not identified'}",
            f"- **Recommendation:** `{item['recommended_review_status'].replace('_', ' ').title()}` — {item['recommendation_rationale']}",
            f"- **Scoring outcome:** `{item['scoring_outcome']}`",
            f"- **ICP:** `{candidate['scoring'].get('score')}` / `{candidate['scoring'].get('band')}`",
            f"- **Confidence:** `{candidate['confidence']['score']}` / `{candidate['confidence']['level']}`",
            f"- **Batch duplicate:** `{candidate['duplicate_check']['batch']['status']}`",
            f"- **HubSpot/Twenty:** `unavailable` / `unavailable`",
            f"- **A2:** `not authorised`",
            "",
            "### Confirmed criteria",
            "",
        ])
        lines.extend([f"- `{criterion}`" for criterion in item["confirmed_criteria"]] or ["- None"])
        lines.extend(["", "### Limitations", ""])
        lines.extend([f"- {limitation}" for limitation in item["limitations"]] or ["- No additional evidence limitation recorded."])
        lines.extend(["", "### Sources", ""])
        lines.extend([f"- {url}" for url in item["source_urls"]])
        lines.extend(["", "### Public contacts", ""])
        lines.extend([
            f"- **{contact.get('label') or contact['contact_type']}**: {contact['value']}"
            for contact in candidate["public_contacts"]
        ] or ["- No public contact published."])
        lines.extend(["", "### Evidence", ""])
        for evidence in candidate["source_evidence"]:
            lines.append(f"- **{evidence['evidence_id']} — {evidence['source_name']}**: {evidence['evidence_excerpt']} ({evidence['source_url']})")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_csv(rows: list[dict[str, Any]], path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REVIEW_COLUMNS, extrasaction="raise")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: "" if row.get(column) is None else row.get(column) for column in REVIEW_COLUMNS})


def style_sheet(sheet, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column[:50]]
        width = min(max(max((len(value) for value in values), default=8) + 2, 12), 60)
        sheet.column_dimensions[column[0].column_letter].width = width


def add_rows(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)


def write_xlsx(package: dict[str, Any], rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    queue = workbook.active
    queue.title = "Review Queue"
    queue_headers = ["Reviewer Decision"] + QUEUE_COLUMNS
    queue.append(queue_headers)
    for row in rows:
        queue.append([""] + [row.get(column) for column in QUEUE_COLUMNS])
    decision_validation = DataValidation(type="list", formula1='"Accept,Reject,Needs Research"', allow_blank=True)
    queue.add_data_validation(decision_validation)
    if queue.max_row >= 2:
        decision_validation.add(f"A2:A{queue.max_row}")
    style_sheet(queue)

    candidate_data = workbook.create_sheet("Candidate Data")
    add_rows(candidate_data, REVIEW_COLUMNS, [[row.get(column) for column in REVIEW_COLUMNS] for row in rows])

    evidence_rows: list[list[Any]] = []
    criteria_rows: list[list[Any]] = []
    score_rows: list[list[Any]] = []
    for item in package["candidates"]:
        candidate = item["prospect_candidate"]
        candidate_id = candidate["candidate_id"]
        for evidence in candidate["source_evidence"]:
            evidence_rows.append([
                candidate_id, evidence["evidence_id"], evidence["source_name"], evidence["source_url"],
                evidence["source_type"], evidence["source_policy_status"], evidence["retrieved_at"],
                evidence["fact_or_inference"], as_json(evidence["supports_claims"]), evidence["evidence_excerpt"],
            ])
        for criterion in candidate["qualification"]["criteria"]:
            criteria_rows.append([
                candidate_id, criterion["criterion_id"], criterion["criterion_label"], criterion["category"],
                criterion["status"], as_json(criterion["evidence_ids"]), criterion.get("notes"),
            ])
        for component in candidate["scoring"]["components"]:
            score_rows.append([
                candidate_id, component["criterion_id"], component["points_awarded"], component["max_points"],
                as_json(component["evidence_ids"]), component.get("notes"),
            ])

    evidence_sheet = workbook.create_sheet("Evidence")
    add_rows(evidence_sheet, ["candidate_id", "evidence_id", "source_name", "source_url", "source_type", "source_policy_status", "retrieved_at", "fact_or_inference", "supports_claims", "evidence_excerpt"], evidence_rows)
    criteria_sheet = workbook.create_sheet("Criteria")
    add_rows(criteria_sheet, ["candidate_id", "criterion_id", "criterion_label", "category", "status", "evidence_ids", "notes"], criteria_rows)
    score_sheet = workbook.create_sheet("Score Components")
    add_rows(score_sheet, ["candidate_id", "criterion_id", "points_awarded", "max_points", "evidence_ids", "notes"], score_rows)

    readme = workbook.create_sheet("Read Me")
    readme_rows = [
        ["Purpose", "Human review of Equinet A1 ICP Discovery candidates."],
        ["Canonical source", "The JSON export is authoritative; this workbook is a derived review view."],
        ["Accept", "Reviewer accepts the candidate for a future authorised next step. It does not trigger A2."],
        ["Reject", "Reviewer rejects the candidate; a reason must be recorded in the authorised workflow."],
        ["Needs Research", "Reviewer requests more evidence or duplicate resolution."],
        ["HubSpot", "Unavailable — customer, opportunity, consent, duplicate and exclusion status are unknown."],
        ["Twenty", "Unavailable — no staging or review action was performed."],
        ["A2", "Not triggered — recorded human approval and an authorised workflow are required."],
        ["Outreach", "Prohibited for A1."],
        ["Package ID", package["package_id"]],
        ["Generated at", package["generated_at"]],
    ]
    add_rows(readme, ["Field", "Meaning"], readme_rows)
    workbook.save(path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def verify_outputs(package: dict[str, Any], json_path: Path, csv_path: Path, xlsx_path: Path) -> None:
    if load_json(json_path) != package:
        raise ValueError("JSON round-trip changed the canonical review package.")
    expected = {item["candidate_id"]: flatten(item) for item in package["candidates"]}
    with csv_path.open("r", encoding="utf-8", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    if [row["candidate_id"] for row in csv_rows] != [item["candidate_id"] for item in package["candidates"]]:
        raise ValueError("CSV candidate order or IDs differ from the review package.")
    for row in csv_rows:
        expected_row = expected[row["candidate_id"]]
        if row["icp_score"] != ("" if expected_row["icp_score"] is None else str(expected_row["icp_score"])):
            raise ValueError(f"CSV score mismatch for {row['candidate_id']}.")
        if row["source_urls"] != expected_row["source_urls"]:
            raise ValueError(f"CSV source URL mismatch for {row['candidate_id']}.")
        if row["full_address"] != (expected_row["full_address"] or ""):
            raise ValueError(f"CSV full-address mismatch for {row['candidate_id']}.")
        if row["public_contact_details"] != expected_row["public_contact_details"]:
            raise ValueError(f"CSV public-contact-detail mismatch for {row['candidate_id']}.")

    workbook = load_workbook(xlsx_path, data_only=False)
    required_sheets = {"Review Queue", "Candidate Data", "Evidence", "Criteria", "Score Components", "Read Me"}
    if not required_sheets.issubset(set(workbook.sheetnames)):
        raise ValueError("Excel workbook is missing a required worksheet.")
    queue = workbook["Review Queue"]
    queue_headers = [cell.value for cell in queue[1]]
    queue_id_column = queue_headers.index("candidate_id") + 1
    queue_ids = [queue.cell(row=index, column=queue_id_column).value for index in range(2, queue.max_row + 1)]
    if queue_ids != [item["candidate_id"] for item in package["candidates"]]:
        raise ValueError("Excel Review Queue order or IDs differ from the review package.")
    candidate_data = workbook["Candidate Data"]
    headers = [cell.value for cell in candidate_data[1]]
    id_column = headers.index("candidate_id") + 1
    score_column = headers.index("icp_score") + 1
    url_column = headers.index("source_urls") + 1
    address_column = headers.index("full_address") + 1
    contact_details_column = headers.index("public_contact_details") + 1
    xlsx_ids = [candidate_data.cell(row=index, column=id_column).value for index in range(2, candidate_data.max_row + 1)]
    if xlsx_ids != [item["candidate_id"] for item in package["candidates"]]:
        raise ValueError("Excel Candidate Data order or IDs differ from the review package.")
    for row_index, candidate_id in enumerate(xlsx_ids, start=2):
        expected_row = expected[candidate_id]
        if candidate_data.cell(row=row_index, column=score_column).value != expected_row["icp_score"]:
            raise ValueError(f"Excel score mismatch for {candidate_id}.")
        if candidate_data.cell(row=row_index, column=url_column).value != expected_row["source_urls"]:
            raise ValueError(f"Excel source URL mismatch for {candidate_id}.")
        if candidate_data.cell(row=row_index, column=address_column).value != expected_row["full_address"]:
            raise ValueError(f"Excel full-address mismatch for {candidate_id}.")
        if candidate_data.cell(row=row_index, column=contact_details_column).value != expected_row["public_contact_details"]:
            raise ValueError(f"Excel public-contact-detail mismatch for {candidate_id}.")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    raise ValueError(f"Unexpected formula in {sheet.title}!{cell.coordinate}.")
                if cell.data_type == "e":
                    raise ValueError(f"Excel error value in {sheet.title}!{cell.coordinate}: {cell.value}")


def export(package_path: Path, output_dir: Path, prefix: str) -> dict[str, Any]:
    package = load_json(package_path)
    validate_schema(REVIEW_SCHEMA, package, "review package")
    for item in package["candidates"]:
        validate_candidate(item["prospect_candidate"])
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "json": output_dir / f"{prefix}.json",
        "markdown": output_dir / f"{prefix}.md",
        "csv": output_dir / f"{prefix}.csv",
        "xlsx": output_dir / f"{prefix}.xlsx",
    }
    rows = [flatten(item) for item in package["candidates"]]
    write_json(package, paths["json"])
    write_markdown(package, paths["markdown"])
    write_csv(rows, paths["csv"])
    write_xlsx(package, rows, paths["xlsx"])
    verify_outputs(package, paths["json"], paths["csv"], paths["xlsx"])

    manifest = {
        "schema_version": "1.0.0",
        "source_package_id": package["package_id"],
        "candidate_count": len(package["candidates"]),
        "files": [
            {"format": format_name, "file_name": path.name, "sha256": sha256(path), "bytes": path.stat().st_size}
            for format_name, path in paths.items()
        ],
        "validation": {
            "review_package": "passed",
            "prospect_candidates": "passed",
            "json_round_trip": "passed",
            "csv_consistency": "passed",
            "xlsx_consistency": "passed",
            "xlsx_formula_errors": "passed",
        },
    }
    validate_schema(MANIFEST_SCHEMA, manifest, "export manifest")
    manifest_path = output_dir / f"{prefix}-manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("review_package", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--prefix", default="equinet-a1-prospects")
    args = parser.parse_args()
    try:
        manifest = export(args.review_package, args.output_dir, args.prefix)
        print(json.dumps(manifest, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(f"INVALID: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
