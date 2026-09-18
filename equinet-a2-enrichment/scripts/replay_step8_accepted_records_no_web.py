#!/usr/bin/env python3
"""Replay accepted Step 8 canonical records and review exports without Web access."""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from a2_wave3_contracts import load
from render_a2_review_views import DEFAULT_SPEC, render_bundle
from validate_a2_enrichment_record import validate
from validate_step2h_review_views import validate_sample

OUT = ROOT / "evaluations/step9/replay"
MANIFEST = OUT / "replay-manifest.json"
CASES = [
    {
        "candidate_id": "A1-DARLEY-JONABELL-001",
        "accepted": ROOT / "evaluations/step8/pilot/A1-DARLEY-JONABELL-001/canonical/revision-005-record-approved.json",
        "previous": ROOT / "evaluations/step8/pilot/A1-DARLEY-JONABELL-001/canonical/revision-004-review-required.json",
        "accepted_review": ROOT / "evaluations/step8/pilot/A1-DARLEY-JONABELL-001/final-review-package",
    },
    {
        "candidate_id": "A1-RR-PODIATRY-001",
        "accepted": ROOT / "evaluations/step8/pilot/A1-RR-PODIATRY-001/canonical/revision-005-held.json",
        "previous": ROOT / "evaluations/step8/pilot/A1-RR-PODIATRY-001/canonical/revision-004-review-required.json",
        "accepted_review": ROOT / "evaluations/step8/pilot/A1-RR-PODIATRY-001/final-review-package",
    },
]
TEXT_EXPORTS = ["review.md", "summary.csv", "review-queue.csv", "fields.csv", "evidence.csv", "requalification.csv", "audit.csv"]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_values(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=True)
    values = {}
    for sheet in workbook.worksheets:
        values[sheet.title] = [[cell.value for cell in row] for row in sheet.iter_rows()]
    workbook.close()
    return values


def main() -> int:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    spec = load(DEFAULT_SPEC)
    results = []
    for case in CASES:
        candidate = case["candidate_id"]
        target = OUT / candidate
        review_dir = target / "review-package"
        target.mkdir(parents=True)
        accepted_record = load(case["accepted"])
        previous_record = load(case["previous"])
        validation = validate(accepted_record, previous=previous_record)
        replay_record = target / "canonical.json"
        replay_record.write_text(json.dumps(accepted_record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        canonical_equal = replay_record.read_bytes() == case["accepted"].read_bytes()
        render_record = case["accepted"].relative_to(ROOT)
        render_previous = case["previous"].relative_to(ROOT)
        render_manifest = render_bundle(render_record, review_dir, render_previous, DEFAULT_SPEC)
        review_validation = validate_sample(
            {"name": f"step9-{candidate}", "record": render_record, "previous": render_previous, "output": review_dir},
            spec,
        )
        text_comparisons = {}
        for name in TEXT_EXPORTS:
            accepted_path = case["accepted_review"] / name
            replay_path = review_dir / name
            text_comparisons[name] = accepted_path.exists() and replay_path.exists() and accepted_path.read_bytes() == replay_path.read_bytes()
        accepted_xlsx = case["accepted_review"] / "review.xlsx"
        replay_xlsx = review_dir / "review.xlsx"
        xlsx_equal = workbook_values(accepted_xlsx) == workbook_values(replay_xlsx)
        accepted_manifest = load(case["accepted_review"] / "manifest.json")
        replay_manifest = load(review_dir / "manifest.json")
        manifest_semantic_equal = (
            accepted_manifest.get("canonical_record_sha256") == replay_manifest.get("canonical_record_sha256")
            and accepted_manifest.get("row_counts") == replay_manifest.get("row_counts")
            and accepted_manifest.get("validation") == replay_manifest.get("validation")
        )
        passed = (
            validation["valid"]
            and canonical_equal
            and review_validation["passed"]
            and all(text_comparisons.values())
            and xlsx_equal
            and manifest_semantic_equal
        )
        results.append(
            {
                "candidate_id": candidate,
                "accepted_record": str(case["accepted"].relative_to(ROOT)),
                "accepted_sha256": sha(case["accepted"]),
                "replay_record": str(replay_record.relative_to(ROOT)),
                "replay_sha256": sha(replay_record),
                "canonical_byte_equal": canonical_equal,
                "canonical_validation": validation,
                "review_validation": review_validation,
                "text_export_equality": text_comparisons,
                "xlsx_cell_value_equality": xlsx_equal,
                "review_manifest_semantic_equality": manifest_semantic_equal,
                "row_counts": render_manifest["row_counts"],
                "passed": passed,
            }
        )
    result = {
        "record_type": "step9_exact_no_web_replay",
        "profile": "equinet-a2-enrichment",
        "release_candidate": "equinet-a2-no-integration-1.0.0-rc.1",
        "status": "pass" if all(item["passed"] for item in results) else "failed",
        "candidate_count": len(results),
        "candidates_passed": sum(item["passed"] for item in results),
        "web_calls": 0,
        "provider_calls": 0,
        "external_actions": 0,
        "results": results,
        "errors": [item["candidate_id"] for item in results if not item["passed"]],
    }
    MANIFEST.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": result["status"], "candidates": f"{result['candidates_passed']}/{result['candidate_count']}", "canonical_exact": all(item["canonical_byte_equal"] for item in results), "exports_exact": all(all(item["text_export_equality"].values()) and item["xlsx_cell_value_equality"] and item["review_manifest_semantic_equality"] for item in results), "web_calls": 0, "external_actions": 0}, indent=2))
    return 0 if result["status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
