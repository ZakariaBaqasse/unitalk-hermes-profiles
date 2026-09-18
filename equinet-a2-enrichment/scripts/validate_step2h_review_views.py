#!/usr/bin/env python3
"""Validate Step 2H review-view specifications and generated sample packages."""

from __future__ import annotations

import csv
import copy
import hashlib
import json
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from render_a2_review_views import build_views, canonical_hash, decode_csv_value, queue_included, render_bundle, render_markdown, safe_csv_value

PROFILE_ROOT = Path(__file__).resolve().parents[1]
SPEC_PATH = PROFILE_ROOT / "foundations" / "contracts" / "canonical" / "a2-review-view-spec-0.1.0.json"
ROOT = PROFILE_ROOT / "evaluations" / "step2h"
OUTPUT = ROOT / "technical-validation.json"
PACKAGE_MANIFEST = ROOT / "step2h-package-manifest.json"
CONTRACT_PATH = PROFILE_ROOT / "foundations" / "A2-REVIEW-VIEW-SPECIFICATION.md"
REVIEW_PATH = ROOT / "A2-REVIEW-VIEW-SPECIFICATION-REVIEW.md"
ACCEPTANCE_RECORD = ROOT / "acceptance-record.json"
CARRIAGE_CORRECTION = ROOT / "post-approval-carriage-return-correction.json"
STEP2I_APPROVAL = PROFILE_ROOT / "evaluations" / "step2i" / "approval-record.json"
LEGACY_PACKAGE_MANIFEST = PROFILE_ROOT / "evaluations" / "step2i" / "pre-promotion-evidence" / "step2h-package-manifest.json"
SCHEMA_PATH = PROFILE_ROOT / "foundations" / "contracts" / "canonical" / "a2-enrichment-record.schema.json"
STEP2G_VALIDATION = PROFILE_ROOT / "evaluations" / "step2g" / "technical-validation.json"
REQUIREMENTS_PATH = PROFILE_ROOT / "foundations" / "contracts" / "requirements-foundation.txt"

SAMPLES = [
    {
        "name": "farrier-review-required",
        "record": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-review-required.json",
        "previous": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-review-required-previous.json",
        "output": ROOT / "samples" / "farrier-review-required",
    },
    {
        "name": "farrier-conflict-held",
        "record": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-conflict-held.json",
        "previous": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-conflict-held-previous.json",
        "output": ROOT / "samples" / "farrier-conflict-held",
    },
    {
        "name": "horse-owner-requalification",
        "record": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "horse-owner-complete-requalification.json",
        "previous": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "horse-owner-complete-requalification-previous.json",
        "output": ROOT / "samples" / "horse-owner-requalification",
    },
    {
        "name": "horse-owner-gap-held",
        "record": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "horse-owner-gap-held.json",
        "previous": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "horse-owner-gap-held-previous.json",
        "output": ROOT / "samples" / "horse-owner-gap-held",
    },
    {
        "name": "farrier-protected-update",
        "record": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-protected-update-awaiting-review.json",
        "previous": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-protected-update-awaiting-review-previous.json",
        "output": ROOT / "samples" / "farrier-protected-update",
    },
    {
        "name": "farrier-application-blocked",
        "record": ROOT / "fixtures" / "farrier-application-blocked.json",
        "previous": PROFILE_ROOT / "evaluations" / "step2f" / "fixtures" / "valid" / "farrier-review-required-previous.json",
        "output": ROOT / "samples" / "farrier-application-blocked",
    },
]


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path: Path) -> tuple[list[str], list[dict]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def csv_expected(value: object) -> str:
    protected = safe_csv_value(value)
    return "" if protected is None else str(protected)


def resolved_path(value: str) -> Path:
    path = Path(value)
    return path.resolve() if path.is_absolute() else (PROFILE_ROOT / path).resolve()


def manifest_validation_errors(sample: dict, manifest: dict, spec: dict, record_hash: str) -> list[str]:
    errors = []
    expected_names = {
        spec["formats"]["markdown"]["filename"],
        spec["formats"]["xlsx"]["filename"],
        *{f"{name.replace('_', '-')}.csv" for name in spec["formats"]["csv"]["files"]},
    }
    output_entries = manifest.get("outputs", [])
    paths = [entry.get("path", "") for entry in output_entries]
    resolved_outputs = [resolved_path(path) for path in paths]
    if len(paths) != len(set(paths)):
        errors.append("manifest contains duplicate output paths")
    if {path.name for path in resolved_outputs} != expected_names or len(resolved_outputs) != len(expected_names):
        errors.append("manifest output set is incomplete or contains unexpected files")
    output_root = sample["output"].resolve()
    if any(path.parent != output_root for path in resolved_outputs):
        errors.append("manifest output path escapes the package directory")
    if resolved_path(manifest.get("canonical_record_path", "")) != sample["record"].resolve():
        errors.append("manifest canonical record path mismatch")
    if manifest.get("canonical_record_sha256") != record_hash:
        errors.append("manifest canonical hash mismatch")
    if resolved_path(manifest.get("previous_record_path", "")) != sample["previous"].resolve():
        errors.append("manifest previous record path mismatch")
    previous = load_json(sample["previous"])
    if manifest.get("previous_record_sha256") != canonical_hash(previous):
        errors.append("manifest previous record hash mismatch")
    if resolved_path(manifest.get("spec_path", "")) != SPEC_PATH.resolve() or manifest.get("spec_sha256") != file_hash(SPEC_PATH):
        errors.append("manifest spec path or hash mismatch")
    renderer_path = PROFILE_ROOT / "scripts" / "render_a2_review_views.py"
    if manifest.get("renderer_sha256") != file_hash(renderer_path):
        errors.append("manifest renderer hash mismatch")
    for entry, path in zip(output_entries, resolved_outputs):
        if not path.exists() or file_hash(path) != entry.get("sha256") or path.stat().st_size != entry.get("bytes"):
            errors.append(f"output hash or size mismatch: {path}")
    return errors


def validate_sample(sample: dict, spec: dict) -> dict:
    errors: list[str] = []
    record = load_json(sample["record"])
    record_hash = canonical_hash(record)
    manifest_path = sample["output"] / "manifest.json"
    manifest = load_json(manifest_path)
    views = build_views(record)

    errors.extend(manifest_validation_errors(sample, manifest, spec, record_hash))
    if manifest["spec_version"] != spec["version"]:
        errors.append("manifest spec version mismatch")
    if manifest["row_counts"] != {name: len(rows) for name, rows in views.items()}:
        errors.append("manifest row counts mismatch")
    csv_rows = {}
    for view_name, columns in spec["formats"]["csv"]["files"].items():
        path = sample["output"] / f"{view_name.replace('_', '-')}.csv"
        headers, rows = read_csv(path)
        csv_rows[view_name] = rows
        if headers != columns:
            errors.append(f"CSV header mismatch: {view_name}")
        if len(rows) != len(views[view_name]):
            errors.append(f"CSV row count mismatch: {view_name}")
        if any(row["canonical_record_sha256"] != record_hash for row in rows):
            errors.append(f"CSV canonical hash mismatch: {view_name}")
        expected_rows = [
            {column: csv_expected(row.get(column, "")) for column in columns}
            for row in views[view_name]
        ]
        if rows != expected_rows:
            errors.append(f"CSV value mismatch: {view_name}")

    expected_queue_ids = [item["field_assessment_id"] for item in record["field_assessments"] if queue_included(item)]
    actual_queue_ids = [row["field_assessment_id"] for row in csv_rows["review_queue"]]
    if set(expected_queue_ids) != set(actual_queue_ids):
        errors.append("review queue membership mismatch")
    field_ids = {row["field_assessment_id"] for row in csv_rows["fields"]}
    if not set(actual_queue_ids).issubset(field_ids):
        errors.append("review queue contains an unknown field assessment")

    combined_evidence_ids = {
        item["evidence_id"] for item in record["source_handoff"]["handoff_snapshot"]["candidate_snapshot"]["source_evidence"]
    } | {item["evidence_id"] for item in record["evidence_registry"]}
    if {row["evidence_id"] for row in csv_rows["evidence"]} != combined_evidence_ids:
        errors.append("evidence projection is not lossless by ID")
    evidence_projection = {row["evidence_id"]: row for row in csv_rows["evidence"]}
    for evidence_id in {
        item["evidence_id"] for item in record["source_handoff"]["handoff_snapshot"]["candidate_snapshot"]["source_evidence"]
    }:
        expected_fields = sorted({
            assessment["field_assessment_id"]
            for assessment in record["field_assessments"]
            if evidence_id in (
                (assessment["baseline"] or {}).get("evidence_ids", [])
                + [ref for observation in assessment["observations"] for ref in observation["evidence_ids"]]
                + (assessment["proposed_resolution"] or {}).get("evidence_ids", [])
            )
        })
        expected_signals = sorted({
            signal["signal_id"] for signal in record["requalification"]["signals"]
            if evidence_id in signal["a1_evidence_ids"]
        })
        projected = evidence_projection[evidence_id]
        if sorted(json.loads(projected["supports_field_assessment_ids_json"])) != expected_fields:
            errors.append(f"A1 evidence field-support mismatch: {evidence_id}")
        if sorted(json.loads(projected["supports_requalification_signal_ids_json"])) != expected_signals:
            errors.append(f"A1 evidence signal-support mismatch: {evidence_id}")

    requalification_count = sum(len(record["requalification"][key]) for key in ["signals", "returns", "score_revision_references"])
    if len(csv_rows["requalification"]) != requalification_count:
        errors.append("requalification projection count mismatch")
    requalification_projection = {row["item_id"]: row for row in csv_rows["requalification"]}
    returns_by_id = {item["return_id"]: item for item in record["requalification"]["returns"]}
    for signal in record["requalification"]["signals"]:
        row = requalification_projection[signal["signal_id"]]
        if sorted(json.loads(row["field_assessment_ids_json"])) != sorted(signal["field_assessment_ids"]):
            errors.append(f"signal field-assessment lineage mismatch: {signal['signal_id']}")
    for revision in record["requalification"]["score_revision_references"]:
        row = requalification_projection[revision["score_revision_id"]]
        source_return = returns_by_id[revision["source_requalification_return_id"]]
        if row["source_requalification_return_id"] != revision["source_requalification_return_id"]:
            errors.append(f"score revision return-link mismatch: {revision['score_revision_id']}")
        if row["result_sha256"] != revision["result_sha256"] or row["signal_snapshot_sha256"] != source_return["signal_snapshot_sha256"]:
            errors.append(f"score revision integrity-hash mismatch: {revision['score_revision_id']}")
        if row["receipt_status"] != (source_return["a1_receipt"] or {}).get("status", ""):
            errors.append(f"score revision receipt mismatch: {revision['score_revision_id']}")

    markdown = (sample["output"] / "review.md").read_text(encoding="utf-8")
    if markdown != render_markdown(record, views):
        errors.append("Markdown projection value mismatch")
    for required in [record_hash, "READ-ONLY", "Review Queue", "Action Boundaries"]:
        if required not in markdown:
            errors.append(f"Markdown missing required marker: {required}")

    workbook_path = sample["output"] / "review.xlsx"
    workbook = load_workbook(workbook_path, data_only=False)
    if workbook.sheetnames != spec["formats"]["xlsx"]["sheets"]:
        errors.append("Excel sheet order mismatch")
    formula_cells = []
    error_cells = []
    non_arial_cells = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells.append(f"{sheet.title}!{cell.coordinate}")
                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"}):
                    error_cells.append(f"{sheet.title}!{cell.coordinate}")
                if cell.value is not None and cell.font.name != "Arial":
                    non_arial_cells.append(f"{sheet.title}!{cell.coordinate}")
    if formula_cells:
        errors.append(f"Excel contains formulas: {formula_cells}")
    if error_cells:
        errors.append(f"Excel contains error values: {error_cells}")
    if non_arial_cells:
        errors.append(f"Excel contains non-Arial cells: {non_arial_cells}")

    sheet_mapping = {
        "Summary": "summary",
        "Review Queue": "review_queue",
        "Fields": "fields",
        "Evidence": "evidence",
        "Requalification": "requalification",
        "Audit": "audit",
    }
    for sheet_name, view_name in sheet_mapping.items():
        sheet = workbook[sheet_name]
        expected_rows = len(views[view_name]) + 1
        if sheet.max_row != expected_rows:
            errors.append(f"Excel row count mismatch: {sheet_name}")
        headers = [cell.value for cell in sheet[1]]
        if headers != spec["formats"]["csv"]["files"][view_name]:
            errors.append(f"Excel header mismatch: {sheet_name}")
        columns = spec["formats"]["csv"]["files"][view_name]
        for row_index, expected_row in enumerate(views[view_name], 2):
            for column_index, column in enumerate(columns, 1):
                actual_value = sheet.cell(row=row_index, column=column_index).value
                expected_value = expected_row.get(column, "")
                if actual_value is None and expected_value == "":
                    actual_value = ""
                if actual_value != expected_value:
                    errors.append(f"Excel value mismatch: {sheet_name}!{get_column_letter(column_index)}{row_index}")

    return {
        "name": sample["name"],
        "canonical_record_sha256": record_hash,
        "row_counts": manifest["row_counts"],
        "errors": errors,
        "passed": not errors,
    }


def main() -> int:
    spec = load_json(SPEC_PATH)
    results = [validate_sample(sample, spec) for sample in SAMPLES]
    failures = [f"sample failed: {result['name']}" for result in results if not result["passed"]]

    negative_checks = {}
    invalid_record = load_json(SAMPLES[0]["record"])
    invalid_record["record_metadata"]["schema_version"] = "9.9.9-invalid"
    with tempfile.TemporaryDirectory(prefix="a2-step2h-") as tmp:
        invalid_path = Path(tmp) / "invalid.json"
        invalid_path.write_text(json.dumps(invalid_record, indent=2) + "\n", encoding="utf-8")
        try:
            render_bundle(invalid_path, Path(tmp) / "invalid-output", SAMPLES[0]["previous"], SPEC_PATH)
            negative_checks["invalid_record_rejected"] = False
        except ValueError:
            negative_checks["invalid_record_rejected"] = True

        incompatible_spec = dict(spec)
        incompatible_spec["canonical_input"] = dict(spec["canonical_input"])
        incompatible_spec["canonical_input"]["schema_version"] = "9.9.9-invalid"
        incompatible_spec_path = Path(tmp) / "incompatible-spec.json"
        incompatible_spec_path.write_text(json.dumps(incompatible_spec, indent=2) + "\n", encoding="utf-8")
        try:
            render_bundle(SAMPLES[0]["record"], Path(tmp) / "incompatible-output", SAMPLES[0]["previous"], incompatible_spec_path)
            negative_checks["incompatible_spec_rejected"] = False
        except ValueError:
            negative_checks["incompatible_spec_rejected"] = True

        source_hash_before = file_hash(SAMPLES[0]["record"])
        render_bundle(SAMPLES[0]["record"], Path(tmp) / "readonly-output", SAMPLES[0]["previous"], SPEC_PATH)
        negative_checks["canonical_source_unchanged"] = source_hash_before == file_hash(SAMPLES[0]["record"])

        base_manifest = load_json(SAMPLES[0]["output"] / "manifest.json")
        base_record_hash = canonical_hash(load_json(SAMPLES[0]["record"]))
        omitted_manifest = copy.deepcopy(base_manifest)
        omitted_manifest["outputs"] = omitted_manifest["outputs"][1:]
        negative_checks["manifest_output_omission_rejected"] = bool(manifest_validation_errors(SAMPLES[0], omitted_manifest, spec, base_record_hash))
        duplicate_manifest = copy.deepcopy(base_manifest)
        duplicate_manifest["outputs"].append(copy.deepcopy(duplicate_manifest["outputs"][0]))
        negative_checks["manifest_duplicate_output_rejected"] = bool(manifest_validation_errors(SAMPLES[0], duplicate_manifest, spec, base_record_hash))
        traversal_manifest = copy.deepcopy(base_manifest)
        traversal_manifest["outputs"][0]["path"] = "../outside-review-package.md"
        negative_checks["manifest_path_traversal_rejected"] = bool(manifest_validation_errors(SAMPLES[0], traversal_manifest, spec, base_record_hash))
        stale_manifest = copy.deepcopy(base_manifest)
        stale_manifest["renderer_sha256"] = "0" * 64
        negative_checks["manifest_stale_renderer_hash_rejected"] = bool(manifest_validation_errors(SAMPLES[0], stale_manifest, spec, base_record_hash))

        corrupted_output = Path(tmp) / "corrupted-package"
        shutil.copytree(SAMPLES[0]["output"], corrupted_output)
        evidence_csv = corrupted_output / "evidence.csv"
        headers, evidence_data = read_csv(evidence_csv)
        evidence_data[0]["summary"] = "Corrupted projection value"
        with evidence_csv.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows(evidence_data)
        corrupted_book_path = corrupted_output / "review.xlsx"
        corrupted_book = load_workbook(corrupted_book_path)
        evidence_sheet = corrupted_book["Evidence"]
        evidence_headers = [cell.value for cell in evidence_sheet[1]]
        evidence_sheet.cell(row=2, column=evidence_headers.index("summary") + 1).value = "Corrupted projection value"
        corrupted_book.save(corrupted_book_path)
        corrupted_manifest_path = corrupted_output / "manifest.json"
        corrupted_manifest = load_json(corrupted_manifest_path)
        for entry in corrupted_manifest["outputs"]:
            path = corrupted_output / Path(entry["path"]).name
            entry["path"] = str(path)
            entry["sha256"] = file_hash(path)
            entry["bytes"] = path.stat().st_size
        corrupted_manifest_path.write_text(json.dumps(corrupted_manifest, indent=2) + "\n", encoding="utf-8")
        corrupted_sample = dict(SAMPLES[0])
        corrupted_sample["output"] = corrupted_output
        corrupted_result = validate_sample(corrupted_sample, spec)
        negative_checks["cross_format_value_corruption_rejected"] = not corrupted_result["passed"]

        injection_record = copy.deepcopy(load_json(SAMPLES[0]["record"]))
        injection_record["subject"]["entities"][0]["display_name"] = "=2+3"
        injection_path = Path(tmp) / "injection-record.json"
        injection_path.write_text(json.dumps(injection_record, indent=2) + "\n", encoding="utf-8")
        injection_output = Path(tmp) / "injection-output"
        render_bundle(injection_path, injection_output, SAMPLES[0]["previous"], SPEC_PATH)
        _, injection_summary = read_csv(injection_output / "summary.csv")
        injection_book = load_workbook(injection_output / "review.xlsx", data_only=False)
        summary_sheet = injection_book["Summary"]
        summary_headers = [cell.value for cell in summary_sheet[1]]
        subject_column = summary_headers.index("subject") + 1
        excel_subject = summary_sheet.cell(row=2, column=subject_column).value
        excel_subject_type = summary_sheet.cell(row=2, column=subject_column).data_type
        decoded_subject = decode_csv_value(injection_summary[0]["subject"])
        negative_checks["spreadsheet_formula_injection_escaped"] = (
            decoded_subject.startswith("=2+3")
            and injection_summary[0]["subject"] != decoded_subject
            and isinstance(excel_subject, str)
            and excel_subject.startswith("=2+3")
            and excel_subject_type == "s"
        )
        reserved_literal = "__A2_TEXT_B64__:literal"
        encoded_literal = safe_csv_value(reserved_literal)
        formula_values = ["=2+3", "+2", "-2", "@cmd", "\tformula", "\rformula", "'=literal", reserved_literal]
        negative_checks["csv_escape_is_reversible_and_collision_free"] = (
            encoded_literal != reserved_literal
            and decode_csv_value(str(encoded_literal)) == reserved_literal
            and encoded_literal != safe_csv_value("=2+3")
            and all(decode_csv_value(str(safe_csv_value(value))) == value for value in formula_values)
        )

        carriage_record = copy.deepcopy(load_json(SAMPLES[0]["record"]))
        carriage_record["subject"]["entities"][0]["display_name"] = "\rformula"
        carriage_path = Path(tmp) / "carriage-record.json"
        carriage_path.write_text(json.dumps(carriage_record, indent=2) + "\n", encoding="utf-8")
        carriage_output = Path(tmp) / "carriage-output"
        render_bundle(carriage_path, carriage_output, SAMPLES[0]["previous"], SPEC_PATH)
        _, carriage_summary = read_csv(carriage_output / "summary.csv")
        carriage_book = load_workbook(carriage_output / "review.xlsx", data_only=False)
        carriage_sheet = carriage_book["Summary"]
        carriage_headers = [cell.value for cell in carriage_sheet[1]]
        carriage_excel = carriage_sheet.cell(row=2, column=carriage_headers.index("subject") + 1).value
        negative_checks["carriage_return_prefix_roundtrip"] = (
            decode_csv_value(carriage_summary[0]["subject"]).startswith("\rformula")
            and isinstance(carriage_excel, str)
            and decode_csv_value(carriage_excel).startswith("\rformula")
        )

    for name, passed in negative_checks.items():
        if not passed:
            failures.append(f"negative check failed: {name}")

    contract_text = CONTRACT_PATH.read_text(encoding="utf-8")
    review_text = REVIEW_PATH.read_text(encoding="utf-8")
    documentation_checks = {
        "contract_version": "**Specification version:** `0.1.0`" in contract_text,
        "contract_approved": "**Step:** `2H — Review-View Specification`" in contract_text and "PROMOTED BY UNITALK OPERATIONS IN STEP 2I" in contract_text,
        "review_version": "**Version:** `0.1.0`" in review_text,
        "decisions_h1_to_h10": all(f"| H{number} |" in review_text for number in range(1, 11)),
        "canonical_json_authority": "only lossless source of truth" in contract_text,
        "step2i_boundary": "Step 2I" in contract_text,
    }
    for name, passed in documentation_checks.items():
        if not passed:
            failures.append(f"documentation check failed: {name}")

    package_paths = [
        SPEC_PATH,
        Path(__file__),
        PROFILE_ROOT / "scripts" / "render_a2_review_views.py",
        SCHEMA_PATH,
        STEP2G_VALIDATION,
        REQUIREMENTS_PATH,
        CONTRACT_PATH,
        REVIEW_PATH,
        ROOT / "INDEPENDENT-ADVERSARIAL-REVIEW-RESOLUTION.md",
    ]
    for sample in SAMPLES:
        package_paths.extend([sample["record"], sample["previous"]])
        package_paths.extend(path for path in sample["output"].iterdir() if path.is_file())
    package_paths = sorted(set(package_paths))
    package_manifest = {
        "manifest_id": "equinet-a2-step2h-review-view-package",
        "version": "0.1.0",
        "status": "approved_by_unitalk_for_step_2i",
        "files": [
            {"path": str(path.relative_to(PROFILE_ROOT)), "sha256": file_hash(path)}
            for path in package_paths
        ],
    }
    PACKAGE_MANIFEST.write_text(json.dumps(package_manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    acceptance = load_json(ACCEPTANCE_RECORD)
    carriage_correction = load_json(CARRIAGE_CORRECTION)
    step2i_approval = load_json(STEP2I_APPROVAL)
    acceptance_checks = {
        "decision": acceptance.get("decision") == "approved_as_drafted",
        "spec_version": acceptance.get("spec_version") == "0.1.0",
        "historical_package_manifest_hash": acceptance.get("package_manifest_sha256") == file_hash(LEGACY_PACKAGE_MANIFEST),
        "carriage_correction_status": carriage_correction.get("status") == "approved_under_h9",
        "historical_carriage_correction_hash": carriage_correction.get("corrected_package_manifest_sha256") == file_hash(LEGACY_PACKAGE_MANIFEST),
        "next_gate": acceptance.get("next_gate") == "Step 2I — Canonical Data Contract Review and Promotion",
        "step2i_promotion_approved": step2i_approval.get("decision") == "approved" and step2i_approval.get("promotion_authorized") is True,
    }
    for name, passed in acceptance_checks.items():
        if not passed:
            failures.append(f"acceptance check failed: {name}")

    result = {
        "step": "2H",
        "spec_version": spec["version"],
        "approval_state": "promoted_in_step_2i",
        "summary": {
            "sample_packages": len(results),
            "sample_packages_passed": sum(result["passed"] for result in results),
            "negative_checks": len(negative_checks),
            "negative_checks_passed": sum(negative_checks.values()),
        },
        "formats": ["markdown", "csv", "xlsx"],
        "samples": results,
        "negative_checks": negative_checks,
        "documentation_checks": documentation_checks,
        "acceptance_checks": acceptance_checks,
        "package_manifest": {"path": str(PACKAGE_MANIFEST), "sha256": file_hash(PACKAGE_MANIFEST), "file_count": len(package_paths)},
        "external_actions": 0,
        "independent_post_fix_review": "pass",
        "failures": failures,
        "pass": not failures,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0 if not failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
