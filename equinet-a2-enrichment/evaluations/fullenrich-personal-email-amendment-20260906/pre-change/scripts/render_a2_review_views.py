#!/usr/bin/env python3
"""Render read-only Markdown, CSV and Excel review views from one canonical A2 record."""

from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from validate_a2_enrichment_record import validate

PROFILE_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SPEC = PROFILE_ROOT / "foundations" / "contracts" / "canonical" / "a2-review-view-spec-0.1.0.json"
RENDERER_VERSION = "0.1.0"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def canonical_bytes(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def canonical_hash(value: object) -> str:
    return hashlib.sha256(canonical_bytes(value)).hexdigest()


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def json_cell(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


CSV_TEXT_PREFIX = "__A2_TEXT_B64__:"
FORMULA_PREFIXES = {"=", "+", "-", "@", "\t", "\r"}


def safe_csv_value(value: object) -> object:
    if isinstance(value, str) and (value[:1] in FORMULA_PREFIXES or value.startswith(CSV_TEXT_PREFIX)):
        encoded = base64.urlsafe_b64encode(value.encode("utf-8")).decode("ascii")
        return CSV_TEXT_PREFIX + encoded
    return value


def decode_csv_value(value: str) -> str:
    if value.startswith(CSV_TEXT_PREFIX):
        return base64.urlsafe_b64decode(value[len(CSV_TEXT_PREFIX):].encode("ascii")).decode("utf-8")
    return value


def safe_excel_value(value: object) -> object:
    if isinstance(value, str) and value[:1] in {"\t", "\r"}:
        return safe_csv_value(value)
    return value


def display(value: object) -> str:
    if value is None:
        return "Unknown"
    if isinstance(value, (dict, list)):
        return json_cell(value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def priority(assessment: dict) -> str:
    proposal = assessment["proposed_resolution"] or {}
    application = assessment["application"]
    if application["status"] in {"blocked", "failed"} or application["workflow_dependency_status"] in {"blocked", "error"} or application["error"] is not None:
        return "blocked"
    if proposal.get("protected_status") in {"protected_manual", "protected_policy", "protected_authoritative"}:
        return "protected"
    if assessment["field_quality_status"] == "conflict" or proposal.get("conflict_status") in {"possible", "material"}:
        return "conflict"
    if assessment["field_quality_status"] == "gap":
        return "gap"
    if proposal.get("action") in {"add", "update", "clear_request", "hold"}:
        return "change"
    return "review"


def queue_included(assessment: dict) -> bool:
    proposal = assessment["proposed_resolution"] or {}
    application = assessment["application"]
    return (
        assessment["field_review"]["decision"] in {"pending", "held", "needs_changes"}
        or assessment["field_quality_status"] in {"gap", "conflict", "partial", "invalid", "error"}
        or proposal.get("action") in {"add", "update", "clear_request", "hold"}
        or proposal.get("conflict_status") in {"possible", "material"}
        or proposal.get("protected_status") in {"protected_manual", "protected_policy", "protected_authoritative"}
        or application["status"] in {"blocked", "failed"}
        or application["workflow_dependency_status"] in {"blocked", "error"}
        or application["error"] is not None
    )


def field_row(record_hash: str, assessment: dict) -> dict:
    baseline = assessment["baseline"] or {}
    proposal = assessment["proposed_resolution"] or {}
    evidence_ids = list(baseline.get("evidence_ids", []))
    for observation in assessment["observations"]:
        evidence_ids.extend(observation["evidence_ids"])
    evidence_ids.extend(proposal.get("evidence_ids", []))
    return {
        "canonical_record_sha256": record_hash,
        "field_assessment_id": assessment["field_assessment_id"],
        "field_key": assessment["field_key"],
        "scope": assessment["scope"],
        "target_id": assessment["entity_id"] or "",
        "presence_status": assessment["presence_status"],
        "availability_status": assessment["availability_status"],
        "field_quality_status": assessment["field_quality_status"],
        "baseline_value_json": json_cell(baseline.get("normalised_value")),
        "observations_json": json_cell(assessment["observations"]),
        "proposed_action": proposal.get("action", ""),
        "proposed_value_json": json_cell(proposal.get("normalised_value")),
        "verification_status": proposal.get("verification_status", ""),
        "confidence_level": proposal.get("confidence_level", ""),
        "freshness_status": proposal.get("freshness_status", ""),
        "conflict_status": proposal.get("conflict_status", ""),
        "protected_status": proposal.get("protected_status", baseline.get("protected_status", "")),
        "evidence_ids_json": json_cell(unique(evidence_ids)),
        "field_review_decision": assessment["field_review"]["decision"],
        "field_review_reason": assessment["field_review"]["reason"] or "",
        "application_status": assessment["application"]["status"],
        "workflow_dependency_status": assessment["application"]["workflow_dependency_status"],
        "application_error_json": json_cell(assessment["application"]["error"]),
        "external_action_reference": assessment["application"]["external_action_reference"] or "",
    }


def queue_row(record_hash: str, assessment: dict) -> dict:
    row = field_row(record_hash, assessment)
    return {
        "canonical_record_sha256": record_hash,
        "priority": priority(assessment),
        "field_assessment_id": row["field_assessment_id"],
        "field_key": row["field_key"],
        "scope": row["scope"],
        "target_id": row["target_id"],
        "baseline_value_json": row["baseline_value_json"],
        "observations_json": row["observations_json"],
        "proposed_action": row["proposed_action"],
        "proposed_value_json": row["proposed_value_json"],
        "verification_status": row["verification_status"],
        "confidence_level": row["confidence_level"],
        "freshness_status": row["freshness_status"],
        "conflict_status": row["conflict_status"],
        "protected_status": row["protected_status"],
        "evidence_ids_json": row["evidence_ids_json"],
        "field_quality_status": row["field_quality_status"],
        "field_review_decision": row["field_review_decision"],
        "application_status": row["application_status"],
        "workflow_dependency_status": row["workflow_dependency_status"],
        "application_error_json": row["application_error_json"],
    }


def evidence_rows(record_hash: str, record: dict) -> list[dict]:
    rows = []
    a1_support_fields: dict[str, list[str]] = {}
    for assessment in record["field_assessments"]:
        references = []
        if assessment["baseline"]:
            references.extend(assessment["baseline"]["evidence_ids"])
        for observation in assessment["observations"]:
            references.extend(observation["evidence_ids"])
        if assessment["proposed_resolution"]:
            references.extend(assessment["proposed_resolution"]["evidence_ids"])
        for evidence_id in references:
            a1_support_fields.setdefault(evidence_id, []).append(assessment["field_assessment_id"])
    a1_support_signals: dict[str, list[str]] = {}
    for signal in record["requalification"]["signals"]:
        for evidence_id in signal["a1_evidence_ids"]:
            a1_support_signals.setdefault(evidence_id, []).append(signal["signal_id"])
    for item in record["source_handoff"]["handoff_snapshot"]["candidate_snapshot"]["source_evidence"]:
        rows.append({
            "canonical_record_sha256": record_hash,
            "evidence_id": item["evidence_id"],
            "namespace": "a1_reference",
            "source_id": item["source_name"],
            "source_type": item["source_type"],
            "source_policy_status": item["source_policy_status"],
            "source_url": item["source_url"],
            "title": item["title"],
            "retrieved_at": item["retrieved_at"],
            "source_date": "",
            "claim_type": item["fact_or_inference"],
            "reliability_level": item["reliability_level"],
            "supports_field_assessment_ids_json": json_cell(unique(a1_support_fields.get(item["evidence_id"], []))),
            "supports_requalification_signal_ids_json": json_cell(unique(a1_support_signals.get(item["evidence_id"], []))),
            "cost_status": "not_applicable",
            "summary": item["evidence_excerpt"],
            "data_minimisation_notes": "Referenced from immutable A1 handoff.",
        })
    for item in record["evidence_registry"]:
        provider_cost = item["provider_cost"] or {}
        rows.append({
            "canonical_record_sha256": record_hash,
            "evidence_id": item["evidence_id"],
            "namespace": item["namespace"],
            "source_id": item["source_id"],
            "source_type": item["source_type"],
            "source_policy_status": item["source_policy_status"],
            "source_url": item["source_url"] or "",
            "title": item["title"] or "",
            "retrieved_at": item["retrieved_at"],
            "source_date": item["source_date"] or "",
            "claim_type": item["claim_type"],
            "reliability_level": item["reliability_level"],
            "supports_field_assessment_ids_json": json_cell(item["supports_field_assessment_ids"]),
            "supports_requalification_signal_ids_json": json_cell(item["supports_requalification_signal_ids"]),
            "cost_status": provider_cost.get("cost_status", "not_applicable"),
            "summary": item["excerpt_or_result_summary"],
            "data_minimisation_notes": item["data_minimisation_notes"] or "",
        })
    return rows


def requalification_rows(record_hash: str, record: dict) -> list[dict]:
    rows = []
    signals_by_id = {item["signal_id"]: item for item in record["requalification"]["signals"]}
    returns_by_id = {item["return_id"]: item for item in record["requalification"]["returns"]}
    for item in record["requalification"]["signals"]:
        rows.append({
            "canonical_record_sha256": record_hash,
            "kind": "signal",
            "item_id": item["signal_id"],
            "status": item["status"],
            "affected_criterion_id": item["affected_criterion_id"],
            "field_assessment_ids_json": json_cell(item["field_assessment_ids"]),
            "materiality": item["materiality"],
            "potential_score_direction": item["potential_score_direction"],
            "source_signal_ids_json": json_cell([item["signal_id"]]),
            "source_requalification_return_id": "",
            "evidence_ids_json": json_cell(item["a1_evidence_ids"] + item["a2_evidence_ids"]),
            "signal_snapshot_sha256": "",
            "result_sha256": "",
            "receipt_status": "",
            "delivery_reference": "",
            "prior_score": "",
            "revised_score": "",
            "prior_band": "",
            "revised_band": "",
            "review_decision": (item["review"] or {}).get("decision", ""),
            "result_reference": "",
        })
    for item in record["requalification"]["returns"]:
        rows.append({
            "canonical_record_sha256": record_hash,
            "kind": "return",
            "item_id": item["return_id"],
            "status": item["status"],
            "affected_criterion_id": "",
            "field_assessment_ids_json": json_cell(unique([field_id for signal in item["signal_snapshots"] for field_id in signal["field_assessment_ids"]])),
            "materiality": "",
            "potential_score_direction": "",
            "source_signal_ids_json": json_cell([signal["signal_id"] for signal in item["signal_snapshots"]]),
            "source_requalification_return_id": "",
            "evidence_ids_json": json_cell(unique([evidence_id for signal in item["signal_snapshots"] for evidence_id in signal["a1_evidence_ids"] + signal["a2_evidence_ids"]])),
            "signal_snapshot_sha256": item["signal_snapshot_sha256"],
            "result_sha256": "",
            "receipt_status": (item["a1_receipt"] or {}).get("status", ""),
            "delivery_reference": item["delivery"]["external_event_reference"] or "",
            "prior_score": "",
            "revised_score": "",
            "prior_band": "",
            "revised_band": "",
            "review_decision": (item["approval"] or {}).get("decision", ""),
            "result_reference": item["delivery"]["external_event_reference"] or "",
        })
    for item in record["requalification"]["score_revision_references"]:
        source_return = returns_by_id.get(item["source_requalification_return_id"], {})
        linked_signals = [signals_by_id[signal_id] for signal_id in item["source_signal_ids"] if signal_id in signals_by_id]
        rows.append({
            "canonical_record_sha256": record_hash,
            "kind": "score_revision",
            "item_id": item["score_revision_id"],
            "status": item["revision_status"],
            "affected_criterion_id": json_cell([change["criterion_id"] for change in item["criteria_changes"]]),
            "field_assessment_ids_json": json_cell(unique([field_id for signal in linked_signals for field_id in signal["field_assessment_ids"]])),
            "materiality": "",
            "potential_score_direction": "",
            "source_signal_ids_json": json_cell(item["source_signal_ids"]),
            "source_requalification_return_id": item["source_requalification_return_id"],
            "evidence_ids_json": json_cell(item["validated_evidence_ids"]),
            "signal_snapshot_sha256": source_return.get("signal_snapshot_sha256", ""),
            "result_sha256": item["result_sha256"],
            "receipt_status": (source_return.get("a1_receipt") or {}).get("status", ""),
            "delivery_reference": (source_return.get("delivery") or {}).get("external_event_reference", ""),
            "prior_score": item["prior_scoring"]["score"],
            "revised_score": item["revised_scoring"]["score"],
            "prior_band": item["prior_scoring"]["band"],
            "revised_band": item["revised_scoring"]["band"],
            "review_decision": (item["review"] or {}).get("decision", ""),
            "result_reference": item["result_reference"],
        })
    return rows


def audit_rows(record_hash: str, record: dict) -> list[dict]:
    rows = []
    for item in record["audit_and_consumption"]["events"]:
        external = item["external_action"] or {}
        rows.append({
            "canonical_record_sha256": record_hash,
            "event_id": item["event_id"],
            "event_type": item["event_type"],
            "timestamp": item["timestamp"],
            "actor_id": item["actor"]["actor_id"],
            "actor_role": item["actor"]["actor_role"],
            "model": item["model"] or "",
            "tools_json": json_cell(item["tools"]),
            "input_references_json": json_cell(item["input_references"]),
            "output_references_json": json_cell(item["output_references"]),
            "external_action_state": external.get("state", ""),
            "external_action_reference": external.get("external_reference", ""),
            "usage_json": json_cell(item["usage"]),
        })
    return rows


def build_views(record: dict) -> dict[str, list[dict]]:
    record_hash = canonical_hash(record)
    candidate = record["source_handoff"]["handoff_snapshot"]["candidate_snapshot"]
    fields = [field_row(record_hash, item) for item in record["field_assessments"]]
    queue = [queue_row(record_hash, item) for item in record["field_assessments"] if queue_included(item)]
    priority_order = {name: index for index, name in enumerate(["blocked", "protected", "conflict", "gap", "change", "review"])}
    queue.sort(key=lambda row: (priority_order[row["priority"]], row["field_key"], row["field_assessment_id"]))
    summary = [{
        "canonical_record_sha256": record_hash,
        "a2_record_id": record["record_metadata"]["a2_record_id"],
        "record_revision_id": record["record_metadata"]["record_revision_id"],
        "segment": candidate["segment"],
        "subject": " | ".join(item["display_name"] for item in record["subject"]["entities"]),
        "workflow_state": record["workflow"]["state"],
        "record_decision": record["review"]["record_decision"],
        "data_quality_status": record["data_quality"]["status"],
        "a2_eligibility_status": record["duplicate_and_eligibility"]["a2_eligibility_status"],
        "a1_score": candidate["scoring"]["score"],
        "a1_band": candidate["scoring"]["band"],
        "field_count": len(fields),
        "review_queue_count": len(queue),
        "evidence_count": len(record["source_handoff"]["handoff_snapshot"]["candidate_snapshot"]["source_evidence"]) + len(record["evidence_registry"]),
        "signal_count": len(record["requalification"]["signals"]),
        "score_revision_count": len(record["requalification"]["score_revision_references"]),
        "limitations_json": json_cell(record["data_quality"]["limitations"] + record["governance"]["policy_limitations"]),
    }]
    return {
        "summary": summary,
        "review_queue": queue,
        "fields": fields,
        "evidence": evidence_rows(record_hash, record),
        "requalification": requalification_rows(record_hash, record),
        "audit": audit_rows(record_hash, record),
    }


def markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    def clean(value: object) -> str:
        return display(value).replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    lines.extend("| " + " | ".join(clean(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def render_markdown(record: dict, views: dict[str, list[dict]]) -> str:
    summary = views["summary"][0]
    queue = views["review_queue"]
    requal = views["requalification"]
    evidence = views["evidence"]
    lines = [
        "# A2 Enrichment Review",
        "",
        f"**Canonical record:** `{summary['a2_record_id']}`  ",
        f"**Revision:** `{summary['record_revision_id']}`  ",
        f"**Canonical SHA-256:** `{summary['canonical_record_sha256']}`  ",
        "**Projection status:** `READ-ONLY — NOT A DECISION OR EXTERNAL ACTION`",
        "",
        "## Decision Summary",
        "",
        markdown_table(
            ["Segment", "Subject", "Workflow", "Record decision", "Data quality", "A2 eligibility", "A1 score/band", "Review items"],
            [[summary["segment"], summary["subject"], summary["workflow_state"], summary["record_decision"], summary["data_quality_status"], summary["a2_eligibility_status"], f"{summary['a1_score']} / {summary['a1_band']}", summary["review_queue_count"]]],
        ),
        "",
        "## Review Queue",
        "",
    ]
    if queue:
        lines.append(markdown_table(
            ["Priority", "Field", "Current", "Observations", "Proposed action", "Proposed", "Evidence", "Confidence", "Freshness", "Conflict", "Protected", "Decision"],
            [[row["priority"], row["field_key"], row["baseline_value_json"], row["observations_json"], row["proposed_action"], row["proposed_value_json"], row["evidence_ids_json"], row["confidence_level"], row["freshness_status"], row["conflict_status"], row["protected_status"], row["field_review_decision"]] for row in queue],
        ))
    else:
        lines.append("No field currently meets the review-queue rules.")
    lines.extend(["", "## Gaps, Conflicts and Limitations", ""])
    quality = record["data_quality"]
    lines.append(markdown_table(
        ["Missing", "Unverified", "Conflicts", "Stale", "Invalid", "Limitations"],
        [[json_cell(quality["missing_field_keys"]), json_cell(quality["unverified_field_keys"]), json_cell(quality["conflict_field_keys"]), json_cell(quality["stale_field_keys"]), json_cell(quality["invalid_field_keys"]), summary["limitations_json"]]],
    ))
    lines.extend(["", "## Requalification", ""])
    if requal:
        lines.append(markdown_table(
            ["Kind", "ID", "Status", "Criterion", "Field assessments", "Direction", "Signals", "Source return", "Prior score", "Revised score", "Receipt", "Review"],
            [[row["kind"], row["item_id"], row["status"], row["affected_criterion_id"], row["field_assessment_ids_json"], row["potential_score_direction"], row["source_signal_ids_json"], row["source_requalification_return_id"], row["prior_score"], row["revised_score"], row["receipt_status"], row["review_decision"]] for row in requal],
        ))
    else:
        lines.append("No requalification item is recorded in this revision.")
    lines.extend(["", "## Evidence Summary", ""])
    lines.append(markdown_table(
        ["ID", "Namespace", "Source", "Policy", "Reliability", "URL", "Summary"],
        [[row["evidence_id"], row["namespace"], row["source_type"], row["source_policy_status"], row["reliability_level"], row["source_url"], row["summary"]] for row in evidence],
    ))
    lines.extend([
        "",
        "## Action Boundaries",
        "",
        "- This view is derived from canonical JSON and is not authoritative data storage.",
        "- It does not approve a field, record, score revision, CRM write or outreach action.",
        "- A1 remains the only numeric score-revision producer.",
        "- External actions require the separately authorised workflow and receipt.",
        "",
    ])
    return "\n".join(lines)


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({column: safe_csv_value(row.get(column, "")) for column in columns} for row in rows)


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column in enumerate(sheet.columns, 1):
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width


def render_xlsx(path: Path, spec: dict, record_path: Path, record_hash: str, views: dict[str, list[dict]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    fixed_time = datetime(2026, 1, 1, 0, 0, 0)
    workbook.properties.creator = "Unitalk AI"
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time

    readme = workbook.create_sheet("Read Me")
    readme_rows = [
        ["A2 Review View", "Read-only deterministic projection"],
        ["Canonical record", str(record_path)],
        ["Canonical SHA-256", record_hash],
        ["Spec version", spec["version"]],
        ["Authority", "Canonical JSON only"],
        ["Reviewer input", "Not supported in this view package"],
        ["External actions", "None"],
        ["A1 scoring", "Displayed only; A2 does not calculate replacement score"],
    ]
    for row in readme_rows:
        readme.append(row)
    readme["A1"].font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    for row in readme.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 90
    readme.sheet_view.showGridLines = False

    mappings = [
        ("Summary", "summary"),
        ("Review Queue", "review_queue"),
        ("Fields", "fields"),
        ("Evidence", "evidence"),
        ("Requalification", "requalification"),
        ("Audit", "audit"),
    ]
    for sheet_name, view_name in mappings:
        sheet = workbook.create_sheet(sheet_name)
        columns = spec["formats"]["csv"]["files"][view_name]
        sheet.append(columns)
        for row in views[view_name]:
            values = [safe_excel_value(row.get(column, "")) for column in columns]
            sheet.append(values)
            row_index = sheet.max_row
            for column_index, value in enumerate(values, 1):
                if isinstance(value, str) and value[:1] in FORMULA_PREFIXES:
                    sheet.cell(row=row_index, column=column_index).data_type = "s"
        style_sheet(sheet)
        if sheet_name == "Evidence":
            url_column = columns.index("source_url") + 1
            for row_index in range(2, sheet.max_row + 1):
                cell = sheet.cell(row=row_index, column=url_column)
                if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    workbook.save(path)


def render_bundle(record_path: Path, output_dir: Path, previous_path: Path | None = None, spec_path: Path = DEFAULT_SPEC) -> dict:
    spec = load_json(spec_path)
    record = load_json(record_path)
    previous = load_json(previous_path) if previous_path else None
    validation = validate(record, previous=previous)
    if not validation["valid"]:
        raise ValueError("Canonical record validation failed: " + " | ".join(validation["errors"]))
    if spec["version"] != "0.1.0" or spec["canonical_input"]["schema_version"] != record["record_metadata"]["schema_version"]:
        raise ValueError("Review-view spec and canonical record versions are incompatible")

    output_dir.mkdir(parents=True, exist_ok=True)
    views = build_views(record)
    record_hash = canonical_hash(record)

    markdown_path = output_dir / spec["formats"]["markdown"]["filename"]
    markdown_path.write_text(render_markdown(record, views), encoding="utf-8")

    csv_paths = {}
    for view_name, columns in spec["formats"]["csv"]["files"].items():
        path = output_dir / f"{view_name.replace('_', '-')}.csv"
        write_csv(path, columns, views[view_name])
        csv_paths[view_name] = path

    xlsx_path = output_dir / spec["formats"]["xlsx"]["filename"]
    render_xlsx(xlsx_path, spec, record_path, record_hash, views)

    output_paths = [markdown_path, *csv_paths.values(), xlsx_path]
    manifest = {
        "spec_version": spec["version"],
        "renderer_version": RENDERER_VERSION,
        "canonical_record_path": str(record_path),
        "canonical_record_sha256": record_hash,
        "previous_record_path": str(previous_path) if previous_path else None,
        "previous_record_sha256": canonical_hash(previous) if previous is not None else None,
        "spec_path": str(spec_path),
        "spec_sha256": file_hash(spec_path),
        "renderer_sha256": file_hash(Path(__file__)),
        "outputs": [
            {"path": str(path), "sha256": file_hash(path), "bytes": path.stat().st_size}
            for path in output_paths
        ],
        "row_counts": {name: len(rows) for name, rows in views.items()},
        "validation": {
            "canonical_record_valid": True,
            "read_only_projection": True,
            "external_actions": 0,
            "formula_policy": "no_formulas",
        },
    }
    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("record", type=Path)
    parser.add_argument("--previous-record", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    args = parser.parse_args()
    try:
        manifest = render_bundle(args.record, args.output_dir, args.previous_record, args.spec)
        print(json.dumps({
            "valid": True,
            "output_dir": str(args.output_dir),
            "canonical_record_sha256": manifest["canonical_record_sha256"],
            "row_counts": manifest["row_counts"],
            "external_actions": 0,
        }, indent=2, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"valid": False, "error": str(exc), "external_actions": 0}, indent=2, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
